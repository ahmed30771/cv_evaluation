"""Minimal text CV for pipeline smoke test."""

from pathlib import Path

from docx import Document

out = Path("uploads") / "smoke_cv.docx"
out.parent.mkdir(exist_ok=True)
doc = Document()
doc.add_heading("Jane Doe", level=1)
doc.add_paragraph("jane.doe@email.com | +1 555 0100")
doc.add_heading("Summary", level=2)
doc.add_paragraph("Software engineer with 4 years building web APIs.")
doc.add_heading("Skills", level=2)
doc.add_paragraph("Python, FastAPI, React, PostgreSQL")
doc.add_heading("Experience", level=2)
doc.add_paragraph("Software Engineer, Acme Corp (2021-Present)")
doc.add_paragraph("- Built billing API serving 20k users")
doc.add_paragraph("- Reduced p95 latency by 35%")
doc.add_heading("Education", level=2)
doc.add_paragraph("B.S. Computer Science, State University, 2021")
doc.save(out)
print(out)
