"""PDF/DOCX text extraction and light section heuristics."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SECTION_HEADERS = {
    "personal_information": [
        r"personal\s+information",
        r"contact",
        r"profile",
    ],
    "summary": [
        r"summary",
        r"professional\s+summary",
        r"objective",
        r"about\s+me",
    ],
    "skills": [
        r"skills",
        r"technical\s+skills",
        r"core\s+competencies",
    ],
    "experience": [
        r"experience",
        r"work\s+experience",
        r"employment",
        r"professional\s+experience",
    ],
    "education": [
        r"education",
        r"academic",
    ],
    "projects": [
        r"projects",
        r"personal\s+projects",
    ],
    "certifications": [
        r"certifications?",
        r"certificates?",
        r"licenses?",
    ],
}


def extract_text_from_file(path: Path, file_type: str) -> str:
    if file_type == "pdf":
        return _extract_pdf(path)
    if file_type == "docx":
        return _extract_docx(path)
    raise ValueError(f"Unsupported file type: {file_type}")


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        parts.append(text)
    return "\n".join(parts).strip()


def _extract_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text).strip()


def parse_sections(raw_text: str) -> dict:
    lines = [ln.strip() for ln in raw_text.splitlines()]
    # Map line index -> section key when a header-like line matches
    header_positions: list[tuple[int, str]] = []
    for i, line in enumerate(lines):
        key = _match_header(line)
        if key:
            header_positions.append((i, key))

    sections: dict = {
        "personal_information": None,
        "summary": None,
        "skills": [],
        "experience": [],
        "education": [],
        "projects": [],
        "certifications": [],
    }

    if not header_positions:
        # Best-effort: first non-empty lines as personal/summary-ish blob
        preview = "\n".join(lines[:12]).strip()
        sections["summary"] = preview or None
        return sections

    # Text before first header → personal_information heuristic
    first_idx = header_positions[0][0]
    preamble = "\n".join(lines[:first_idx]).strip()
    if preamble:
        sections["personal_information"] = {"raw": preamble}

    for idx, (start, key) in enumerate(header_positions):
        end = header_positions[idx + 1][0] if idx + 1 < len(header_positions) else len(lines)
        body_lines = [ln for ln in lines[start + 1 : end] if ln]
        body = "\n".join(body_lines).strip()
        if key == "skills":
            skills = re.split(r"[,|•·\n]", body)
            sections["skills"] = [s.strip() for s in skills if s.strip()]
        elif key == "summary":
            sections["summary"] = body or None
        elif key == "personal_information":
            sections["personal_information"] = {"raw": body}
        elif key in ("experience", "education", "projects", "certifications"):
            bullets = [ln.lstrip("•-* ").strip() for ln in body_lines if ln]
            sections[key] = [{"raw": b} for b in bullets] if bullets else []
        else:
            sections[key] = body

    return sections


def _match_header(line: str) -> str | None:
    cleaned = line.strip().lower().rstrip(":")
    if len(cleaned) > 48:
        return None
    for key, patterns in SECTION_HEADERS.items():
        for pat in patterns:
            if re.fullmatch(pat, cleaned):
                return key
    return None


def extract_from_bytes(data: bytes, file_type: str) -> str:
    if file_type == "pdf":
        reader = PdfReader(BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages).strip()
    if file_type == "docx":
        doc = Document(BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text).strip()
    raise ValueError(f"Unsupported file type: {file_type}")
